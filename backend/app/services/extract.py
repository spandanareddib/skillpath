import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader


class ExtractionError(Exception):
    """Raised when a resume file cannot be turned into usable text."""


def _clean(text: str) -> str:
    text = text.replace("\r", "\n").replace("\x00", "")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _read_pdf(path: Path) -> str:
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise ExtractionError("PDF is password protected")
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("Could not read this PDF") from exc


def _read_docx(path: Path) -> str:
    try:
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    except Exception as exc:
        raise ExtractionError("Could not read this DOCX") from exc


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        raw = _read_pdf(path)
    elif ext == ".docx":
        raw = _read_docx(path)
    else:
        raise ExtractionError("Unsupported file type")
    text = _clean(raw)
    if not text:
        raise ExtractionError("No extractable text found (is this a scanned PDF?)")
    return text
